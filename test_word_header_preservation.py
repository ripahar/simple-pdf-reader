import fitz
from docx import Document
from PIL import Image, ImageDraw
from word_header_preservation import preserve_repeated_pdf_header


def test_preserves_repeated_header_when_top_regions_match(tmp_path):
    pdf_path = tmp_path / "repeat.pdf"
    pdf = fitz.open()
    for i in range(3):
        page = pdf.new_page(width=612, height=792)
        page.draw_rect(fitz.Rect(0, 0, 612, 70), color=(0, 0, 0), fill=(0.1, 0.25, 0.5))
        page.insert_text((30, 35), "GENERIC HEADER", color=(1, 1, 1), fontsize=14)
        page.insert_text((30, 120), f"Body page {i+1}")
    pdf.save(pdf_path)
    pdf.close()

    docx_path = tmp_path / "out.docx"
    doc = Document()
    doc.add_paragraph("Editable body")
    doc.save(docx_path)

    result = preserve_repeated_pdf_header(str(pdf_path), str(docx_path))
    assert result["preserved"] is True
    reopened = Document(docx_path)
    assert len(reopened.sections[0].header.paragraphs[0].runs[0]._r.xpath(".//w:drawing")) > 0


def test_skips_unique_page_tops(tmp_path):
    pdf_path = tmp_path / "unique.pdf"
    pdf = fitz.open()
    for i in range(3):
        page = pdf.new_page(width=612, height=792)
        page.draw_rect(fitz.Rect(0, 0, 612, 70), color=(0, 0, 0), fill=(i * 0.25, 0.1, 0.2))
        page.insert_text((30, 120), f"Body page {i+1}")
    pdf.save(pdf_path)
    pdf.close()

    docx_path = tmp_path / "out.docx"
    doc = Document()
    doc.add_paragraph("Editable body")
    doc.save(docx_path)

    result = preserve_repeated_pdf_header(str(pdf_path), str(docx_path))
    assert result["preserved"] is False
