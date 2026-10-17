"""Generic header/banner preservation for PDF-to-DOCX conversion.

The helper detects a repeated visual strip at the top of pages and preserves it
as a Word header image. It is brand-agnostic and does not contain document-
specific names, colors, or coordinates. It deliberately avoids injecting a
header when page-top regions are not sufficiently similar.
"""
from __future__ import annotations

import io
from pathlib import Path


def _render_top_band(page, fraction=0.105, scale=2.0):
    import fitz
    rect = page.rect
    clip = fitz.Rect(rect.x0, rect.y0, rect.x1, rect.y0 + rect.height * fraction)
    pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), clip=clip, alpha=False)
    from PIL import Image
    return Image.open(io.BytesIO(pix.tobytes("png"))).convert("RGB")


def _similarity(image_a, image_b):
    from PIL import ImageChops, ImageStat
    if image_a.size != image_b.size:
        image_b = image_b.resize(image_a.size)
    diff = ImageChops.difference(image_a, image_b)
    mean_diff = sum(ImageStat.Stat(diff).mean) / 3.0
    return 1.0 - mean_diff / 255.0


def preserve_repeated_pdf_header(pdf_path: str, docx_path: str) -> dict:
    """Preserve repeated page-top artwork as a Word header image if confidently detected."""
    import fitz
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    pdf = fitz.open(pdf_path)
    try:
        if len(pdf) < 2:
            return {"preserved": False, "reason": "PDF has fewer than two pages."}
        indices = sorted(set([0, len(pdf) // 2, len(pdf) - 1]))
        bands = [_render_top_band(pdf[i]) for i in indices]
        scores = [_similarity(bands[0], band) for band in bands[1:]]
        if not scores or min(scores) < 0.94:
            return {"preserved": False, "reason": "Page-top regions were not sufficiently similar; no header image inserted.", "similarity_scores": scores}
        header_png = Path(docx_path).with_name(Path(docx_path).stem + "_detected_header.png")
        bands[0].save(header_png, format="PNG")
        doc = Document(docx_path)
        for section in doc.sections:
            header = section.header
            p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            if not p._p.xpath(".//w:drawing"):
                p.add_run().add_picture(str(header_png), width=Inches(7.4))
        doc.save(docx_path)
        try:
            header_png.unlink()
        except OSError:
            pass
        return {"preserved": True, "reason": "Repeated top-of-page visual strip inserted into Word header.", "similarity_scores": scores}
    finally:
        pdf.close()
