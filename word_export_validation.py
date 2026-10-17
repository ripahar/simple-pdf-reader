"""Post-conversion checks for PDF-to-DOCX exports.

This module intentionally reports possible defects; it does not silently modify
Word content or claim that pixel-perfect layout can be validated without a
Word-compatible renderer.
"""
from __future__ import annotations

import re
from collections import Counter
from pathlib import Path


def _normalize(text: str) -> str:
    text = text.replace("\u00ad", "")
    text = re.sub(r"(?<=\w)-\s+(?=\w)", "", text)  # rejoin line-break hyphenation
    return re.sub(r"\s+", " ", text).strip().casefold()


def _pdf_text(path: str) -> tuple[str, int]:
    import fitz
    doc = fitz.open(path)
    try:
        return "\n".join(page.get_text("text") for page in doc), len(doc)
    finally:
        doc.close()


def _docx_text(path: str) -> tuple[str, int, int, int]:
    from docx import Document
    doc = Document(path)
    chunks = [p.text for p in doc.paragraphs]
    table_count = len(doc.tables)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                chunks.append(cell.text)
    for section in doc.sections:
        chunks.extend(p.text for p in section.header.paragraphs)
        chunks.extend(p.text for p in section.footer.paragraphs)
    return "\n".join(chunks), len(doc.sections), table_count, len(doc.inline_shapes)


def validate_pdf_to_docx(pdf_path: str, docx_path: str) -> dict:
    """Return metrics and warnings for a converted PDF/DOCX pair."""
    source_text, source_pages = _pdf_text(pdf_path)
    output_text, section_count, table_count, image_count = _docx_text(docx_path)
    source_norm = _normalize(source_text)
    output_norm = _normalize(output_text)
    token_re = re.compile(r"[^\W_]+(?:['’][^\W_]+)?", re.UNICODE)
    source_tokens = token_re.findall(source_norm)
    output_tokens = token_re.findall(output_norm)
    source_counts = Counter(source_tokens)
    output_counts = Counter(output_tokens)
    total_source_tokens = sum(source_counts.values())
    matched_tokens = sum(min(count, output_counts[token]) for token, count in source_counts.items())
    coverage = matched_tokens / total_source_tokens if total_source_tokens else 1.0
    missing_counts = source_counts - output_counts
    missing_sample = [word for word, count in missing_counts.most_common(18)]
    warnings = []
    if coverage < 0.98:
        warnings.append(f"Possible text loss: only {coverage:.1%} of source PDF word occurrences were found in the DOCX text.")
    if source_pages != section_count:
        warnings.append(f"PDF has {source_pages} pages while DOCX has {section_count} sections. DOCX section count is not the same as rendered page count; use a Word/LibreOffice renderer for an exact page-count comparison.")
    if not table_count:
        warnings.append("No editable Word tables were detected.")
    if not image_count:
        warnings.append("No inline images were detected; verify graphics and logos manually.")
    return {"source_pdf": str(Path(pdf_path)), "output_docx": str(Path(docx_path)), "source_pages": source_pages, "docx_sections_not_page_count": section_count, "docx_tables": table_count, "docx_inline_images": image_count, "source_word_occurrences": total_source_tokens, "matched_word_occurrences": matched_tokens, "text_coverage": coverage, "missing_word_sample": missing_sample, "warnings": warnings}


def format_validation_report(result: dict) -> str:
    lines = ["PDF TO WORD CONVERSION VALIDATION", "=" * 36, f"Source PDF: {result['source_pdf']}", f"Output DOCX: {result['output_docx']}", f"Source PDF pages: {result['source_pages']}", f"DOCX sections (NOT rendered page count): {result['docx_sections_not_page_count']}", f"Editable Word tables detected: {result['docx_tables']}", f"Inline images detected: {result['docx_inline_images']}", f"Source word occurrences: {result['source_word_occurrences']}", f"Matched word occurrences: {result['matched_word_occurrences']}", f"Text occurrence coverage: {result['text_coverage']:.1%}", "", "Possible missing words (sample): " + (", ".join(result['missing_word_sample']) or "None detected"), "", "WARNINGS:"]
    if result['warnings']:
        lines.extend(f"- {warning}" for warning in result['warnings'])
    else:
        lines.append("- No automatic text-coverage warnings. Visual review is still required.")
    lines += ["", "Limitations:", "- This report detects suspicious text loss; it does not repair the DOCX.", "- Token coverage can be affected by OCR errors, ligatures, or unusual reading order.", "- Exact page count and visual overlap checks require rendering with Microsoft Word or LibreOffice."]
    return "\n".join(lines) + "\n"
