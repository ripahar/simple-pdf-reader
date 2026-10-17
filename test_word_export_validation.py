import fitz
from docx import Document

from word_export_validation import validate_pdf_to_docx, format_validation_report


def test_validation_detects_missing_text_and_reports_structure(tmp_path):
    pdf_path = tmp_path / "source.pdf"
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Important rubric criterion score and comments")
    pdf.save(pdf_path)
    pdf.close()

    docx_path = tmp_path / "output.docx"
    doc = Document()
    doc.add_paragraph("Important rubric criterion")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "comments"
    doc.save(docx_path)

    result = validate_pdf_to_docx(str(pdf_path), str(docx_path))
    assert result["text_coverage"] < 1.0
    assert result["docx_tables"] == 1
    assert any("Possible text loss" in warning for warning in result["warnings"])
    assert "Text occurrence coverage" in format_validation_report(result)


def test_validation_handles_complete_text(tmp_path):
    pdf_path = tmp_path / "source.pdf"
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "A complete sentence for validation.")
    pdf.save(pdf_path)
    pdf.close()

    docx_path = tmp_path / "output.docx"
    doc = Document()
    doc.add_paragraph("A complete sentence for validation.")
    doc.save(docx_path)

    result = validate_pdf_to_docx(str(pdf_path), str(docx_path))
    assert result["text_coverage"] == 1.0
    assert not any("Possible text loss" in warning for warning in result["warnings"])
